import os
import re
import sys
import datetime
import pandas as pd
import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font
from azure.storage.blob import BlobServiceClient
from dotenv import load_dotenv

# === LOG SETUP ===
LOG_PATH = f"rapport_log_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.txt"

class Tee:
    def __init__(self, *streams):
        self.streams = streams
    def write(self, msg):
        for s in self.streams:
            s.write(msg); s.flush()
    def flush(self):
        for s in self.streams:
            s.flush()

log_file = open(LOG_PATH, "w", encoding="utf-8")
sys.stdout = Tee(sys.__stdout__, log_file)
sys.stderr = Tee(sys.__stderr__, log_file)

print(f"=== Script start: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')} ===")
print(f"Log gemmes i: {LOG_PATH}\n")

load_dotenv()
ACCOUNT_NAME       = os.getenv("AZURE_STORAGE_ACCOUNT_NAME")
SAS_TOKEN          = os.getenv("AZURE_SAS_TOKEN")
BLOB_PREFIX        = ""
LOCAL_PDF_DIR      = "downloaded_pdfs"
EXCEL_PATH         = "faktura_list.xlsx"
RESULT_PATH        = "sammenligningsrapport_finance.xlsx"
FAKTURA_NR_FIELD   = "faktura nr"
EKSTRA_FAKTURA_DIR = r"C:\Users\mojoh\OneDrive - Semler Gruppen A S\Dokumenter\hist faktura\00032"

assert ACCOUNT_NAME and SAS_TOKEN, "Tjek .env!"

FIELDS = [("Kunde nr", r"Kunde[^\d]*(\d+)")]
os.makedirs(LOCAL_PDF_DIR, exist_ok=True)

DATE_REGEX = r"(\d{1,2}[.\-/]\d{1,2}[.\-/]\d{2,4}|\d{4}[.\-/]\d{1,2}[.\-/]\d{1,2}|\b\d{6,8}\b)"
# Beløb-regex med optionel negativ fortegn
BELOB_REGEX = r"(-?[0-9]{1,3}(?:[.,][0-9]{3})*[.,][0-9]+|-?[0-9]+[.,][0-9]+)"

errors = []
warnings = []

def log_error(msg):
    errors.append(msg); print(f"❌ ERROR: {msg}")

def log_warning(msg):
    warnings.append(msg); print(f"⚠️  WARNING: {msg}")

def pdf_matches_faktura(fname, faktura_nr):
    if not fname.lower().endswith(".pdf"):
        return False
    if faktura_nr in fname:
        return True
    if len(faktura_nr) >= 13:
        del_kode = faktura_nr[4:7]
        del_nr   = faktura_nr[7:]
        if f"{del_kode}_{del_nr}" in fname: return True
        if f"{del_kode}-{del_nr}" in fname: return True
        if f"{del_kode}{del_nr}" in fname: return True
    return False

def safe_str(val):
    if pd.isna(val): return ""
    if isinstance(val, float):
        if val == int(val): return str(int(val))
        return str(val)
    if isinstance(val, int): return str(val)
    if isinstance(val, (pd.Timestamp, datetime.datetime, datetime.date)):
        return val.strftime("%Y-%m-%d")
    s = str(val).strip()
    m = re.match(r"(\d{4}-\d{2}-\d{2})\s+\d{1,2}:\d{2}", s)
    if m: return m.group(1)
    if re.fullmatch(r"\d{4}-\d{2}-\d{2} 00:00:00", s): return s[:10]
    return s

def normalize_date_any(val):
    if not val or not isinstance(val, str): return ""
    val = val.strip()
    m = re.match(r"(\d{4}-\d{2}-\d{2})\s+\d{1,2}:\d{2}", val)
    if m: val = m.group(1)
    m = re.fullmatch(r"(\d{4})[.\-/](\d{1,2})[.\-/](\d{1,2})", val)
    if m:
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if 1980 <= y <= 2050 and 1 <= mo <= 12 and 1 <= d <= 31:
            try: return datetime.date(y, mo, d).strftime("%Y-%m-%d")
            except ValueError: pass
    m = re.fullmatch(r"(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{4})", val)
    if m:
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if 1980 <= y <= 2050 and 1 <= mo <= 12 and 1 <= d <= 31:
            try: return datetime.date(y, mo, d).strftime("%Y-%m-%d")
            except ValueError: pass
    m = re.fullmatch(r"(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{2})", val)
    if m:
        d, mo = int(m.group(1)), int(m.group(2)); y = int(m.group(3))
        y = y + 2000 if y < 80 else y + 1900
        if 1 <= mo <= 12 and 1 <= d <= 31:
            try: return datetime.date(y, mo, d).strftime("%Y-%m-%d")
            except ValueError: pass
    m = re.fullmatch(r"(\d{8})", val)
    if m:
        s = m.group(1)
        y, mo, d = int(s[:4]), int(s[4:6]), int(s[6:8])
        if 1980 <= y <= 2050 and 1 <= mo <= 12 and 1 <= d <= 31:
            try: return datetime.date(y, mo, d).strftime("%Y-%m-%d")
            except ValueError: pass
        d, mo, y = int(s[:2]), int(s[2:4]), int(s[4:8])
        if 1980 <= y <= 2050 and 1 <= mo <= 12 and 1 <= d <= 31:
            try: return datetime.date(y, mo, d).strftime("%Y-%m-%d")
            except ValueError: pass
    return ""

def normalize_date_alt(val):
    if not val or not isinstance(val, str): return ""
    val = val.strip()
    m = re.match(r"(\d{4}-\d{2}-\d{2})\s+\d{1,2}:\d{2}", val)
    if m: val = m.group(1)
    m = re.fullmatch(r"(\d{4})[.\-/](\d{1,2})[.\-/](\d{1,2})", val)
    if m:
        y, d, mo = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if 1980 <= y <= 2050 and 1 <= mo <= 12 and 1 <= d <= 31:
            try: return datetime.date(y, mo, d).strftime("%Y-%m-%d")
            except ValueError: pass
    return ""

def dates_match(ve, vp):
    norm_ve = normalize_date_any(ve)
    norm_vp = normalize_date_any(vp)
    if norm_ve and norm_vp and norm_ve == norm_vp:
        if ve != vp:
            return True, "Samme dato, forskellig formatering"
        return True, ""
    alt_ve = normalize_date_alt(ve)
    if alt_ve and norm_vp and alt_ve == norm_vp:
        return True, f"Excel-dato i YYYY-DD-MM format ('{ve}' = {alt_ve})"
    alt_vp = normalize_date_alt(vp)
    if norm_ve and alt_vp and norm_ve == alt_vp:
        return True, f"PDF-dato i YYYY-DD-MM format ('{vp}' = {alt_vp})"
    return False, ""

def extract_fakturadato(text):
    lines = text.split("\n")
    for line in lines:
        label_match = re.search(r"fakturadato", line, re.IGNORECASE)
        if label_match:
            after_label = line[label_match.end():]
            for m in re.finditer(DATE_REGEX, after_label):
                norm = normalize_date_any(m.group(1))
                if norm: return norm
    for idx, line in enumerate(lines):
        label_match = re.search(r"fakturadato", line, re.IGNORECASE)
        if label_match:
            label_start = label_match.start()
            for offset in range(1, 3):
                if idx + offset < len(lines):
                    nline = lines[idx + offset]
                    if not nline.strip(): continue
                    best_match = None; best_dist = 9999
                    for m in re.finditer(DATE_REGEX, nline):
                        dist = abs(m.start() - label_start)
                        if dist < best_dist:
                            best_dist = dist; best_match = m.group(1)
                    if best_match:
                        norm = normalize_date_any(best_match)
                        if norm: return norm
            break
    return ""

def dansk_til_float_str(raw):
    """Konverterer dansk talformat ('200.000,00' eller '-200.000,00') til standard ('200000.00')."""
    raw = raw.strip()
    neg = raw.startswith("-")
    if neg: raw = raw[1:].strip()
    if "," in raw:
        # Dansk: punktum = tusindsep, komma = decimal
        raw = raw.replace(".", "").replace(",", ".")
    result = ("-" if neg else "") + raw
    return result

def is_credit_nota(text):
    """Tjekker om PDF er en kreditnota."""
    return bool(re.search(r"Kreditnota", text, re.IGNORECASE))

def extract_total_ialt_from_pdf(pdf):
    full_text = "\n".join((page.extract_text() or "") for page in pdf.pages)
    credit = is_credit_nota(full_text)
    
    # Prøv flere mønstre i prioriteret rækkefølge
    patterns = [
        r"14 dage[^\d\-]*" + BELOB_REGEX,
        r"I\s*alt[^\d\-]*" + BELOB_REGEX,
        r"(?:Total|At betale|Fakturabel.b)[^\d\-]*" + BELOB_REGEX,
    ]
    
    result = ""
    # For "I alt": find sidste forekomst (typisk det endelige beløb)
    matches = re.findall(r"I\s*alt[^\d\-]*" + BELOB_REGEX, full_text)
    if matches:
        result = matches[-1].replace(" ", "")
    else:
        for pat in patterns:
            m = re.search(pat, full_text)
            if m:
                result = m.group(1).replace(" ", "")
                break
    
    if not result:
        return ""
    
    result = dansk_til_float_str(result)
    
    # Hvis kreditnota og beløbet er positivt, gør det negativt
    if credit and result and not result.startswith("-"):
        try:
            if float(result) > 0:
                result = "-" + result
        except ValueError:
            pass
    
    return result

def extract_pdf_moms(text):
    """Find moms-beløb i PDF-tekst. Håndterer flere formater og kreditnotaer."""
    credit = is_credit_nota(text)
    moms = ""
    lines = text.split("\n")
    for i, l in enumerate(lines):
        # Skip linjer med negativ-kontekst
        if re.search(r"\b(?:INCL\.?|INKL\.?)\s*MOMS\b", l, re.IGNORECASE):
            continue
        if re.search(r"\b(?:uden|for|ikke|ekskl\.?|exkl\.?)\s+moms", l, re.IGNORECASE):
            continue

        # Format 1: "Salgsmoms 25,00% af 1500,00  375,00"
        m = re.search(r"(?:Salgsmoms|Moms)[^0-9\-]*[0-9.,]+\s*%\s*af\s*-?[0-9.,]+\s+(-?[0-9.,]+)", l, re.IGNORECASE)
        if m:
            moms = dansk_til_float_str(m.group(1))
            continue

        # Format 2: "(HERAF MOMS 4.560,00)" eller "(HERAF MOMS -4.560,00)"
        m = re.search(r"HERAF\s*MOMS[^\d\-]*([\-]?[0-9][0-9.,]*)", l, re.IGNORECASE)
        if m:
            raw = m.group(1).rstrip(")").strip()
            moms = dansk_til_float_str(raw)
            continue

        # Format 3: Linje STARTER med "Moms" eller "Salgsmoms" + beløb (uden procent)
        m = re.match(r"^\s*(?:Salgsmoms|Moms)\s+([\-]?[0-9][0-9.,]*)\s*$", l, re.IGNORECASE)
        if m:
            moms = dansk_til_float_str(m.group(1))
            continue

        # Format 4: Linje er KUN "Moms" eller "Salgsmoms" (label alene), beløb på næste linje
        if re.match(r"^\s*(?:Salgsmoms|Moms)\s*$", l, re.IGNORECASE):
            for offset in range(1, 3):
                if i + offset < len(lines):
                    next_line = lines[i + offset].strip()
                    m2 = re.match(r"^([\-]?[0-9][0-9.,]*)\s*$", next_line)
                    if m2:
                        moms = dansk_til_float_str(m2.group(1))
                        break
    
    # Hvis kreditnota og moms er positiv, gør den negativ
    if credit and moms and not moms.startswith("-"):
        try:
            if float(moms) > 0:
                moms = "-" + moms
        except ValueError:
            pass
    
    return moms

def normalize_number(val):
    if not val or str(val).strip() == "": return "0"
    s = str(val).replace(",", ".").replace(" ", "")
    try: return str(float(s))
    except: return s.lstrip("0")

def floats_equal(v1, v2, tol=0.02):
    try:
        return abs(float(str(v1).replace(",",".")) - float(str(v2).replace(",","."))) < tol
    except: return False

def empty_or_zero(val):
    if val is None or str(val).strip() == "": return True
    try: return float(str(val).replace(",",".").replace(" ","")) == 0.0
    except: return False

def kunde_nrs_equal(k1, k2):
    def num(s):
        return str(int(re.sub(r"\D","",s))) if re.search(r"\d",str(s)) else str(s).strip()
    return num(k1) == num(k2)

def row_score(row):
    score = 0
    for col in ["moms", "i alt", "kunde nr", "faktura dato"]:
        v = row.get(col)
        if pd.isna(v): continue
        try:
            if float(v) != 0: score += 1
        except (ValueError, TypeError):
            if str(v).strip(): score += 1
    return score

# === LÆS EXCEL OG FJERN DUBLETTER ===
print("=== LÆSER EXCEL ===")
df = pd.read_excel(EXCEL_PATH, engine='openpyxl')
df.columns = df.columns.str.strip().str.lower()
print(f"Excel kolonner: {df.columns.tolist()}")
print(f"Rækker FØR dedup: {len(df)}")

df[FAKTURA_NR_FIELD] = df[FAKTURA_NR_FIELD].apply(
    lambda x: str(int(x)) if pd.notna(x) and isinstance(x, (float, int)) and float(x) == int(float(x))
              else str(x).strip() if pd.notna(x) else ""
)

print("\n=== DUBLET-RENSNING ===")
df["_score"] = df.apply(row_score, axis=1)
df = df.sort_values(by=[FAKTURA_NR_FIELD, "_score"], ascending=[True, False])

dup_mask = df.duplicated(subset=[FAKTURA_NR_FIELD], keep=False)
if dup_mask.any():
    duplicates = df[dup_mask]
    unique_dup_nrs = duplicates[FAKTURA_NR_FIELD].unique()
    log_warning(f"Fandt {len(unique_dup_nrs)} faktura-numre med dubletter")
    for nr in unique_dup_nrs:
        sub = duplicates[duplicates[FAKTURA_NR_FIELD] == nr]
        print(f"  {nr}: {len(sub)} rækker")
        for _, r in sub.iterrows():
            kept = "✅ BEHOLDT" if r["_score"] == sub["_score"].max() else "❌ FJERNET"
            print(f"    {kept} | score={r['_score']} | moms={r.get('moms')} | i alt={r.get('i alt')} | kunde nr={r.get('kunde nr')} | dato={r.get('faktura dato')}")
else:
    print("✅ Ingen dubletter fundet")

df = df.drop_duplicates(subset=[FAKTURA_NR_FIELD], keep="first").drop(columns=["_score"])
df = df[df[FAKTURA_NR_FIELD].astype(str).str.strip() != ""]
print(f"Rækker EFTER dedup: {len(df)}")
print("=== END DUBLET-RENSNING ===\n")

faktura_nrs = set(str(val).strip() for val in df[FAKTURA_NR_FIELD] if pd.notna(val) and str(val).strip())

# === AZURE DOWNLOAD (multi-container baseret på dealer-kode) ===
print("=== AZURE DOWNLOAD ===")
BLOB_URL = f"https://{ACCOUNT_NAME}.blob.core.windows.net"
print("Forbinder til Azure Blob Storage:", BLOB_URL)
svc = BlobServiceClient(account_url=BLOB_URL, credential=SAS_TOKEN)

dealer_koder = set()
for fnr in faktura_nrs:
    if len(fnr) >= 13:
        dealer_koder.add(fnr[4:7])
print(f"Fundne dealer-koder: {sorted(dealer_koder)}")

container_navne = [f"00{kode}" for kode in sorted(dealer_koder)]
print(f"Vil scanne containers: {container_navne}")

allerede_hentet = set()
downloaded = 0
total_blobs_scanned = 0

for container_navn in container_navne:
    try:
        az_container = svc.get_container_client(container_navn)
        if not az_container.exists():
            log_warning(f"Container '{container_navn}' findes ikke – springer over")
            continue
        print(f"\nScanner container: {container_navn}")
        container_blobs = 0
        for blob in az_container.list_blobs(name_starts_with=BLOB_PREFIX):
            container_blobs += 1
            total_blobs_scanned += 1
            fn = os.path.basename(blob.name)
            if not fn.lower().endswith(".pdf"):
                continue
            for faktura_nr in faktura_nrs:
                if faktura_nr in allerede_hentet:
                    continue
                if pdf_matches_faktura(fn, faktura_nr):
                    allerede_hentet.add(faktura_nr)
                    blob_client = az_container.get_blob_client(blob)
                    local_path = os.path.join(LOCAL_PDF_DIR, fn)
                    if os.path.exists(local_path):
                        print(f"  Allerede hentet: {fn} → {faktura_nr}")
                    else:
                        try:
                            with open(local_path, "wb") as f:
                                f.write(blob_client.download_blob().readall())
                                print(f"  Hentet: {fn} → {faktura_nr}")
                        except Exception as e:
                            log_error(f"Kunne ikke hente {fn}: {e}")
                    downloaded += 1
                    break
        print(f"  Container {container_navn}: {container_blobs} blobs scannet")
    except Exception as e:
        log_error(f"Fejl ved container '{container_navn}': {type(e).__name__}: {str(e)[:200]}")

print(f"\nTotalt scannede blobs: {total_blobs_scanned}")
print(f"Azure download færdig – {downloaded} PDF'er matched")

# === LOKALE PDF'ER ===
print("\n=== LOKALE PDF'ER ===")
lokal_hentet = set()
if os.path.isdir(EKSTRA_FAKTURA_DIR):
    for fname in os.listdir(EKSTRA_FAKTURA_DIR):
        if not fname.lower().endswith(".pdf"): continue
        for faktura_nr in faktura_nrs:
            if faktura_nr in allerede_hentet or faktura_nr in lokal_hentet: continue
            if pdf_matches_faktura(fname, faktura_nr):
                lokal_hentet.add(faktura_nr)
                src = os.path.join(EKSTRA_FAKTURA_DIR, fname)
                dst = os.path.join(LOCAL_PDF_DIR, fname)
                if not os.path.exists(dst):
                    try:
                        with open(src,"rb") as fsrc, open(dst,"wb") as fdst:
                            fdst.write(fsrc.read())
                        print(f"Kopieret lokal PDF: {fname} → {faktura_nr}")
                    except Exception as e:
                        log_error(f"Fejl ved kopi af {fname}: {e}")
                break
else:
    log_warning(f"Lokal folder findes ikke: {EKSTRA_FAKTURA_DIR}")

# === PARSE PDF'ER ===
print("\n=== PARSER PDF'ER ===")
factura_dict = {}
for _, row in df.iterrows():
    nr = safe_str(row.get(FAKTURA_NR_FIELD, ""))
    if nr and nr.lower() != "nan":
        factura_dict[nr] = row

pdf_data = {}
manglende_pdf = []; manglende_dato = []; manglende_ialt = []; manglende_moms = []
antal_kreditnotaer = 0
alle_filer = os.listdir(LOCAL_PDF_DIR)

for idx, faktura_nr in enumerate(faktura_nrs):
    filnavn = None
    for fname in alle_filer:
        if pdf_matches_faktura(fname, faktura_nr):
            filnavn = fname; break
    if not filnavn:
        manglende_pdf.append(faktura_nr)
        print(f"{idx+1:04d}/{len(faktura_nrs)} ❌ Ingen PDF matcher {faktura_nr}")
        continue
    pdf_path = os.path.join(LOCAL_PDF_DIR, filnavn)
    print(f"{idx+1:04d}/{len(faktura_nrs)} | {filnavn}")
    try:
        with pdfplumber.open(pdf_path) as pdf:
            text = "\n".join((page.extract_text() or "") for page in pdf.pages)
            er_kreditnota = is_credit_nota(text)
            if er_kreditnota:
                antal_kreditnotaer += 1
                print(f"\t📋 Kreditnota detekteret")
            entry = {"PDF fil": filnavn, "Faktura nr": faktura_nr, "Kreditnota": er_kreditnota}
            for col, regex in FIELDS:
                m = re.search(regex, text)
                entry[col] = m.group(1) if m else ""
            dato = extract_fakturadato(text)
            entry["Faktura dato"] = dato
            entry["moms"] = extract_pdf_moms(text)
            entry["i alt"] = extract_total_ialt_from_pdf(pdf)
            pdf_data[faktura_nr] = entry
            if not dato:
                manglende_dato.append(faktura_nr)
                print(f"\t⚠️  Ingen fakturadato fundet!")
            else:
                print(f"\t✅ Fakturadato: {dato}")
            if entry["i alt"]:
                print(f"\t✅ I alt: {entry['i alt']}")
            else:
                manglende_ialt.append(faktura_nr)
                print(f"\t⚠️  Intet 'i alt' beløb fundet!")
            if entry["moms"]:
                print(f"\t✅ Moms: {entry['moms']}")
            else:
                manglende_moms.append(faktura_nr)
                print(f"\t⚠️  Ingen moms fundet!")
    except Exception as e:
        log_error(f"Kunne ikke parse {filnavn}: {e}")
        continue

# === EXCEL RAPPORT ===
print("\n=== BYGGER EXCEL RAPPORT ===")
wb = Workbook()
ws = wb.active
ws.title = "Sammenligning"
header = ["Faktura nr","PDF fil","Type","Kunde nr Excel","Kunde nr PDF",
          "Faktura dato Excel","Faktura dato PDF",
          "moms Excel","moms PDF","i alt Excel","i alt PDF","Kommentar"]
ws.append(header)
red_fill   = PatternFill(start_color="FFC7CE",end_color="FFC7CE",fill_type="solid")
green_fill = PatternFill(start_color="C6EFCE",end_color="C6EFCE",fill_type="solid")
yellow_fill = PatternFill(start_color="FFEB9C",end_color="FFEB9C",fill_type="solid")
bold_font  = Font(bold=True)
red_font   = Font(color="9C0006")
green_font = Font(color="006100")
for cell in ws[1]: cell.font = bold_font

antal_match_ok = 0; antal_match_med_forskel = 0; antal_uden_pdf = 0

for faktura_nr, row in factura_dict.items():
    ex_fnr  = safe_str(row.get("faktura nr",""))
    ex_knr  = safe_str(row.get("kunde nr",""))
    ex_dato = safe_str(row.get("faktura dato",""))
    ex_moms = safe_str(row.get("moms",""))
    ex_ialt = safe_str(row.get("i alt",""))
    pe = pdf_data.get(faktura_nr)
    kommentar = []; style_info = []
    if not pe:
        rr = [ex_fnr,"","",ex_knr,"",ex_dato,"",ex_moms,"",ex_ialt,"",
              "Match ikke fundet i dokumenthotel"]
        ws.append(rr)
        c = ws.cell(row=ws.max_row, column=len(rr))
        c.fill = red_fill; c.font = red_font
        antal_uden_pdf += 1
        continue
    pf = pe.get("PDF fil",""); pk = safe_str(pe.get("Kunde nr",""))
    pd2 = safe_str(pe.get("Faktura dato","")); pm = safe_str(pe.get("moms",""))
    pi  = safe_str(pe.get("i alt",""))
    type_str = "Kreditnota" if pe.get("Kreditnota") else "Faktura"
    rr = [ex_fnr,pf,type_str,ex_knr,pk,ex_dato,pd2,ex_moms,pm,ex_ialt,pi,""]
    # Justér kolonneindex (vi har tilføjet "Type" som kolonne 3)
    pairs = [(3,4,"Kunde nr"),(5,6,"Faktura dato"),(7,8,"moms"),(9,10,"i alt")]
    har_forskel = False
    for ie,ip,label in pairs:
        ve = rr[ie]; vp = rr[ip]; ig = False
        if label == "Kunde nr":
            if empty_or_zero(ve) and empty_or_zero(vp): ig = True
            elif kunde_nrs_equal(ve,vp):
                if str(ve)!=str(vp):
                    kommentar.append(f"{label}: Samme værdi, forskellig formatering")
                ig = True
        elif label == "Faktura dato":
            if empty_or_zero(ve) and empty_or_zero(vp): ig = True
            elif ve and vp:
                match, kom = dates_match(ve, vp)
                if match:
                    if kom:
                        kommentar.append(f"{label}: {kom}")
                    ig = True
        else:
            if empty_or_zero(ve) and empty_or_zero(vp): ig = True
            elif floats_equal(ve,vp):
                if normalize_number(ve)!=normalize_number(vp):
                    kommentar.append(f"{label}: Samme værdi, forskellig formatering")
                ig = True
        if ig:
            style_info.append(("green",ie)); style_info.append(("green",ip))
        else:
            style_info.append(("red",ie)); style_info.append(("red",ip))
            kommentar.append(f"{label}: Forskel (Excel: '{ve}' | PDF: '{vp}')")
            har_forskel = True
    rr[-1] = "; ".join([c for c in kommentar if 'Forskel' in c or 'formatering' in c or 'YYYY-DD-MM' in c])
    ws.append(rr)
    rn = ws.max_row
    # Gul baggrund i Type-kolonnen hvis kreditnota
    if pe.get("Kreditnota"):
        ws.cell(row=rn, column=3).fill = yellow_fill
        ws.cell(row=rn, column=3).font = bold_font
    for style,col in style_info:
        cell = ws.cell(row=rn, column=col+1)
        if style=="red": cell.fill=red_fill; cell.font=red_font
        elif style=="green": cell.fill=green_fill; cell.font=green_font
    if har_forskel: antal_match_med_forskel += 1
    else: antal_match_ok += 1

wb.save(RESULT_PATH)
print(f"Excel-rapport gemt som: {RESULT_PATH}")

# === OPSUMMERING ===
print("\n" + "="*60)
print("=== OPSUMMERING ===")
print("="*60)
print(f"Total fakturaer i Excel:           {len(factura_dict)}")
print(f"  ✅ Match uden forskelle:         {antal_match_ok}")
print(f"  🟡 Match MED forskelle:          {antal_match_med_forskel}")
print(f"  ❌ Uden matchende PDF:           {antal_uden_pdf}")
print(f"  📋 Heraf kreditnotaer:           {antal_kreditnotaer}")
print(f"\nAzure containers scannet:          {len(container_navne)}")
print(f"Azure blobs scannet:               {total_blobs_scanned}")
print(f"PDF'er hentet fra Azure:           {downloaded}")
print(f"PDF'er kopieret lokalt:            {len(lokal_hentet)}")
print(f"\nPDF'er uden fakturadato:           {len(manglende_dato)}")
for nr in manglende_dato: print(f"    - {nr}")
print(f"\nPDF'er uden 'i alt' beløb:         {len(manglende_ialt)}")
for nr in manglende_ialt: print(f"    - {nr}")
print(f"\nPDF'er uden moms:                  {len(manglende_moms)}")
for nr in manglende_moms: print(f"    - {nr}")
print(f"\nFakturaer uden PDF:                {len(manglende_pdf)}")
for nr in manglende_pdf: print(f"    - {nr}")
print(f"\nAntal warnings:                    {len(warnings)}")
for w in warnings: print(f"    ⚠️  {w}")
print(f"\nAntal errors:                      {len(errors)}")
for e in errors: print(f"    ❌ {e}")
print("="*60)
print(f"=== Script slut: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')} ===")
print(f"\nLog gemt i: {LOG_PATH}")

log_file.close()
